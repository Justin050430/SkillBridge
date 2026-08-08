import os
import pdfplumber
from docx import Document
from openpyxl import load_workbook

BASE = r"c:\Projects\SkillBridge\Assignment"
OUT = r"c:\Projects\SkillBridge\Assignment\extracted"

os.makedirs(OUT, exist_ok=True)

def extract_pdf(path, out_name):
    with pdfplumber.open(path) as pdf:
        text = []
        for i, page in enumerate(pdf.pages):
            t = page.extract_text() or ""
            text.append(f"--- PAGE {i+1} ---\n{t}")
        with open(os.path.join(OUT, out_name), "w", encoding="utf-8") as f:
            f.write("\n\n".join(text))
    print(f"PDF done: {out_name}")

def extract_docx(path, out_name):
    doc = Document(path)
    lines = []
    # paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            lines.append(para.text)
    # tables
    for ti, table in enumerate(doc.tables):
        lines.append(f"\n=== TABLE {ti+1} ===")
        for row in table.rows:
            cells = [c.text.replace("\n", " | ") for c in row.cells]
            lines.append("\t".join(cells))
    with open(os.path.join(OUT, out_name), "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"DOCX done: {out_name}")

def extract_xlsx(path, out_name):
    wb = load_workbook(path, data_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"\n=== SHEET: {ws.title} (dims {ws.dimensions}) ===")
        for row in ws.iter_rows():
            vals = [str(c.value) if c.value is not None else "" for c in row]
            if any(vals):
                lines.append("\t".join(vals))
    with open(os.path.join(OUT, out_name), "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"XLSX done: {out_name}")

extract_pdf(os.path.join(BASE, "BMIT3173 Assignment Specification.pdf"), "specification.txt")
extract_pdf(os.path.join(BASE, "Assignment Planning - Sheet1.pdf"), "planning.txt")
extract_docx(os.path.join(BASE, "BMIT3173 Assignment Report.docx"), "report.txt")
extract_xlsx(os.path.join(BASE, "BMIT3173 Assignment Rubrics.xlsx"), "rubrics.txt")

print("ALL DONE")
